import math
import os
import re
import secrets
import datetime
from io import BytesIO

import openpyxl
from flask import Flask, flash, redirect, render_template, request, send_file, url_for

app = Flask(__name__)

# Production-д SECRET_KEY-г Environment табаас өгнө. Санамсаргүй түлхүүр бол
# сервер дахин асах бүрд сессүүд (flash мессеж) тасардаг.
app.secret_key = os.environ.get('SECRET_KEY') or secrets.token_hex(16)

# Ажлын хавтас (CWD) хаана ч байсан эх файлууд олдохын тулд app.py-гийн хавтаснаас тооцно.
BASE_DIR = os.path.dirname(os.path.abspath(__file__))

DEFAULT_FILES = {
    'substitutions': os.path.join(BASE_DIR, 'Орлуулах.xlsx'),
    'required': os.path.join(BASE_DIR, 'ЗБНТЖ.xlsx'),
}

STATE = {'substitutions': [], 'required': [], 'groups': [], 'required_codes': set()}

DEFAULT_MTIMES = {}

CODE_RE = re.compile(r'^[A-ZА-ЯӨҮЁ0-9][A-ZА-ЯӨҮЁ0-9./_-]{2,}$', re.IGNORECASE)
DIGIT_RE = re.compile(r'^\d{6,8}$')
NUM_CLEAN_RE = re.compile(r'[,₮% ]')
DATE_RE = re.compile(r'(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})')


def clean(value):
    if value is None:
        return ''
    if isinstance(value, (datetime.datetime, datetime.date)):
        return value.strftime('%Y-%m-%d')
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def code_like(value):
    return bool(CODE_RE.match(clean(value)))


def get(row, index):
    if 0 <= index < len(row):
        return row[index]
    return ''


def to_number(value):
    s = NUM_CLEAN_RE.sub('', clean(value))
    if s == '':
        return 0.0
    try:
        return float(s)
    except ValueError:
        return float('nan')


def js_number_or_zero(value):
    s = clean(value)
    if s == '':
        return 0
    try:
        n = float(s)
    except ValueError:
        return 0
    return 0 if n == 0 else n


def numeric(row):
    values = [n for n in (to_number(v) for v in row) if math.isfinite(n) and n > 0]
    return values[-1] if values else 0


def sales_amount(row):
    n = to_number(get(row, 19))
    return 0 if math.isnan(n) or n == 0 else n


def rows_from_workbook(data):
    workbook = openpyxl.load_workbook(BytesIO(data), data_only=True, read_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    return [[clean(value) for value in row] for row in sheet.iter_rows(values_only=True)]


def parse_groups(rows):
    header = next((row for row in rows if any('дотоод код' in clean(v).lower() for v in row)), None)
    header_index = rows.index(header) if header is not None else -1
    source_code_index = -1
    alternative_code_indexes = []
    if header is not None:
        source_code_index = next((i for i, v in enumerate(header) if clean(v).lower() == 'дотоод код'), -1)
        alternative_code_indexes = [
            i for i, v in enumerate(header)
            if 'орлуулах' in clean(v).lower() and 'дотоод код' in clean(v).lower()
        ]

    groups = []
    for row_index, row in enumerate(rows):
        if source_code_index >= 0:
            code_index = source_code_index
        else:
            code_index = next((i for i, v in enumerate(row) if code_like(v)), -1)
        if code_index < 0:
            continue

        source = {'code': get(row, code_index), 'name': get(row, code_index + 1) or get(row, code_index)}
        alternatives = []
        if alternative_code_indexes:
            candidate_indexes = alternative_code_indexes
        else:
            length = max(0, math.ceil((len(row) - code_index - 2) / 2))
            candidate_indexes = [code_index + 2 + i * 2 for i in range(length)]
        for index in candidate_indexes:
            value = get(row, index)
            if code_like(value):
                alternatives.append({'code': value, 'name': get(row, index + 1) or value})

        if row_index > header_index and source_code_index >= 0 and code_like(source['code']):
            groups.append([source] + alternatives)

    return groups


def expand_products(groups):
    expanded = []
    for group in groups:
        for product in group:
            expanded.append({
                'code': product['code'],
                'name': product['name'],
                'alternatives': [candidate for candidate in group if candidate['code'] != product['code']],
            })

    product_by_code = {}
    alternatives_by_code = {}
    for product in expanded:
        product_by_code[product['code']] = {'code': product['code'], 'name': product['name']}
        alternatives_by_code.setdefault(product['code'], {})

    for product in expanded:
        for alternative in product['alternatives']:
            alternatives_by_code[product['code']][alternative['code']] = alternative
            alternatives_by_code.setdefault(alternative['code'], {})
            alternatives_by_code[alternative['code']][product['code']] = product_by_code[product['code']]

    return [
        {
            'code': code,
            'name': product['name'],
            'alternatives': list(alternatives_by_code.get(code, {}).values()),
        }
        for code, product in product_by_code.items()
    ]


def parse_products(rows):
    return expand_products(parse_groups(rows))


def header_code_column(rows):
    """'Дотоод код' баганын мөр/багана индексийг олно."""
    for row_index, row in enumerate(rows):
        for column, value in enumerate(row):
            if clean(value).lower() == 'дотоод код':
                return row_index, column
    return -1, -1


def stock_columns(rows, header_index):
    """Толгой хэсгээс 'Үлдэгдэл' баганануудыг олно."""
    columns = []
    for row in rows[header_index:header_index + 3]:
        for column, value in enumerate(row):
            if clean(value).lower() == 'үлдэгдэл' and column not in columns:
                columns.append(column)
    return columns


def required_codes_from(rows):
    """ЗБНТЖ файлын 'Дотоод код' баганаас нэр төрлийн жагсаалт гаргана."""
    header_index, code_column = header_code_column(rows)
    if code_column < 0:
        return set()
    return {
        code for code in (clean(get(row, code_column)) for row in rows[header_index + 1:])
        if code and code_like(code)
    }


def warehouse_stock_codes(rows):
    """Агуулахын тайлангаас үлдэгдэлтэй (>0) нэр төрлийн кодуудыг гаргана.

    Тайланд үлдэгдэлгүй бараа ч жагсдаг тул зөвхөн жагсаалтад байгаагаар нь
    тооцвол хангалт үргэлж 100% гарна.
    """
    header_index, code_column = header_code_column(rows)
    if code_column < 0:
        return set()
    columns = stock_columns(rows, header_index)

    codes = set()
    for row in rows[header_index + 1:]:
        code = clean(get(row, code_column))
        if not code or not code_like(code):
            continue
        if not columns:
            codes.add(code)
            continue
        for column in columns:
            value = to_number(get(row, column))
            if math.isfinite(value) and value > 0:
                codes.add(code)
                break
    return codes


def outage_events(rows):
    """Тасалдлын тайлангаас бүтээгдэхүүн бүрийн алдсан борлуулалтыг нэгтгэнэ.

    Нэг бүтээгдэхүүн хэд хэдэн удаа тасарч болох тул мөр бүрийг тусад нь биш,
    дотоод кодоор нь нэгтгэж тооцно. Тайлангийн 'Нийт' мөр код агуулаагүй тул
    өөрөө шүүгдэж үлдэнэ.
      11 багана = 'Алдсан борлуулалт' (тоо ширхэг)
      12 багана = 'Нийт алдагдсан борлуулалт' (үнийн дүн)
    """
    totals = {}
    names = {}

    def add(code, row):
        entry = totals.setdefault(code, {'units': 0, 'revenue': 0})
        entry['units'] += js_number_or_zero(get(row, 11))
        entry['revenue'] += js_number_or_zero(get(row, 12))
        if code not in names:
            names[code] = ''

    header_index, code_column = header_code_column(rows)
    if code_column >= 0:
        for row in rows[header_index + 1:]:
            code = clean(get(row, code_column))
            if code and code_like(code):
                add(code, row)
                if not names[code]:
                    names[code] = clean(get(row, code_column + 1))
    else:
        # 'Дотоод код' толгой олдоогүй бол хуучин эвристикт шилжинэ.
        for row in rows:
            code = find_code(row)
            if code:
                add(code, row)
                if not names[code]:
                    names[code] = find_product_name(row, code)

    return totals, names


def find_code(row):
    for value in row:
        if DIGIT_RE.match(clean(value)):
            return clean(value)
    for value in row:
        if code_like(value):
            return clean(value)
    return ''


def find_product_name(row, code):
    try:
        index = row.index(code)
    except ValueError:
        return ''
    return get(row, index + 1) or ''


def analysis_days(rows):
    text = ' '.join(clean(value) for row in rows for value in row)
    dates = []
    for match in DATE_RE.finditer(text):
        try:
            dates.append(datetime.date(int(match.group(1)), int(match.group(2)), int(match.group(3))))
        except ValueError:
            continue
    if len(dates) >= 2:
        return max(1, (max(dates) - min(dates)).days + 1)
    days = [n for n in (to_number(get(row, 7)) for row in rows) if math.isfinite(n) and n > 0]
    return max(days) if days else 30


def resolve_file(filename):
    """Файлын нэрийг том/жижиг үсэг харгалзахгүйгээр олно."""
    if os.path.isfile(filename):
        return filename
    folder = os.path.dirname(filename) or '.'
    target = os.path.basename(filename).lower()
    try:
        for entry in os.listdir(folder):
            if entry.lower() == target:
                return os.path.join(folder, entry)
    except OSError:
        pass
    return filename


def load_defaults(force=False):
    """Эх файлуудыг (Орлуулах, ЗБНТЖ) уншина.

    Диск дээрх файл өөрчлөгдсөн үед л дахин уншдаг тул эх файлаа сольсны дараа
    серверийг дахин ажиллуулах шаардлагагүй. Өөрчлөгдсөн бол True буцаана.
    """
    changed = False
    for key, filename in DEFAULT_FILES.items():
        path = resolve_file(filename)
        try:
            mtime = os.path.getmtime(path)
        except OSError:
            mtime = None
        if not force and DEFAULT_MTIMES.get(key) == mtime and STATE[key]:
            continue

        try:
            with open(path, 'rb') as handle:
                rows = rows_from_workbook(handle.read())
            groups = parse_groups(rows)
            STATE[key] = expand_products(groups)
            if key == 'substitutions':
                STATE['groups'] = groups
            if key == 'required':
                STATE['required_codes'] = required_codes_from(rows)
        except Exception:
            STATE[key] = []
            if key == 'substitutions':
                STATE['groups'] = []
            if key == 'required':
                STATE['required_codes'] = set()

        DEFAULT_MTIMES[key] = mtime
        changed = True
    return changed


def analyze(sales_rows, previous_rows, warehouse_rows, outage_rows):
    # ЗБНТЖ-н хангалт: агуулахад үлдэгдэлтэй байгаа ЗБНТЖ нэр төрлийн тоо / ЗБНТЖ-н нийт тоо
    required_codes = STATE['required_codes'] or {item['code'] for item in STATE['required']}
    warehouse_codes = warehouse_stock_codes(warehouse_rows)
    matched_required_codes = required_codes & warehouse_codes
    coverage = min(100, round(len(matched_required_codes) / len(required_codes) * 100)) if required_codes else 0

    outage_map, outage_names = outage_events(outage_rows)
    outage_codes = list(outage_map.keys())

    sale_map = {}
    for row in sales_rows:
        code = find_code(row)
        if code:
            sale_map[code] = {'units': numeric(row), 'revenue': sales_amount(row)}

    previous_map = {}
    for row in previous_rows:
        code = find_code(row)
        if code:
            previous_map[code] = numeric(row)

    report_days = analysis_days(outage_rows)
    substitutions_by_code = {item['code']: item for item in STATE['substitutions']}

    results = []
    for code in outage_codes[:200]:
        outage_name = outage_names.get(code, '')
        source = substitutions_by_code.get(code)
        source = dict(source) if source else {'code': code, 'name': outage_name or 'Тодорхойгүй нэр төрөл', 'alternatives': []}
        if outage_name:
            source['name'] = outage_name
        alternatives = source['alternatives'] if source['alternatives'] else [{'code': '', 'name': 'Орлуулах мэдээлэл байхгүй'}]

        lost_units = outage_map[code]['units']
        lost_revenue = outage_map[code]['revenue']

        alternative_sales = [
            {
                'code': item['code'],
                'name': item['name'],
                'units': sale_map.get(item['code'], {}).get('units', 0),
                'revenue': sale_map.get(item['code'], {}).get('revenue', 0),
            }
            for item in alternatives
        ]
        sold = sum(item['units'] for item in alternative_sales)
        revenue = sum(item['revenue'] for item in alternative_sales)
        daily = sum(previous_map.get(item['code'], 0) for item in alternatives) / 30
        expected_units = daily * report_days
        average_unit_price = revenue / sold if sold > 0 else 0
        excess = max(0, sold - expected_units)
        excess_revenue = max(0, revenue - expected_units * average_unit_price)

        results.append({
            'code': code,
            'name': source['name'],
            'alternatives': alternative_sales,
            'sold': sold,
            'revenue': revenue,
            'lostUnits': lost_units,
            'lostRevenue': lost_revenue,
            'excess': excess,
            'excessRevenue': excess_revenue,
            'required': code in required_codes,
        })

    results.sort(key=lambda row: (-row['sold'], -row['revenue']))
    return {
        'coverage': coverage,
        'matched': len(matched_required_codes),
        'required_total': len(required_codes),
        # Нийт дүнг харуулах 200 мөрийн хязгаараас үл хамааран бүтнээр нь тооцно.
        'lost_units_total': sum(item['units'] for item in outage_map.values()),
        'lost_revenue_total': sum(item['revenue'] for item in outage_map.values()),
        'outage_total': len(outage_map),
        'analysis': results,
    }


def match_rank(item, query):
    """Хайлтын тэмдэгтүүд яг тухайн дарааллаараа хаана таарснаар эрэмбэлнэ.

    Бага тоо = эхэнд. Өөрийнх нь код/нэр таарсан нь орлуулах хувилбараараа
    таарснаас үргэлж түрүүлнэ.
    """
    code = item['code'].lower()
    name = item['name'].lower()
    if code.startswith(query):
        return 0
    if name.startswith(query):
        return 1
    if query in code:
        return 2
    if query in name:
        return 3
    for alt in item['alternatives']:
        if alt['code'].lower().startswith(query) or alt['name'].lower().startswith(query):
            return 4
    for alt in item['alternatives']:
        if query in alt['code'].lower() or query in alt['name'].lower():
            return 5
    return None


def search_substitutions(query):
    query = (query or '').strip().lower()
    if not query:
        return []

    ranked = []
    for item in STATE['substitutions']:
        rank = match_rank(item, query)
        if rank is not None:
            ranked.append((rank, item))
    ranked.sort(key=lambda pair: pair[0])
    return [item for _, item in ranked[:12]]


def money(value):
    return f"₮{round(value or 0):,}"


app.jinja_env.filters['money'] = money

EMPTY_ANALYSIS = {
    'coverage': 0, 'analysis': [], 'matched': 0, 'required_total': 0,
    'lost_units_total': 0, 'lost_revenue_total': 0, 'outage_total': 0,
}

LAST_ANALYSIS = dict(EMPTY_ANALYSIS)

LOGO_FILE = 'monos-logo.png'


def logo_exists():
    return os.path.isfile(os.path.join(app.static_folder, LOGO_FILE))

DOWNLOAD_HEADER = [
    'Дотоод код', 'Тасалдсан нэр төрөл', 'Алдсан борлуулалтын тоо хэмжээ', 'Алдсан борлуулалтын үнийн дүн',
    'Орлуулсан бүтээгдэхүүн', 'Орлуулсан борлуулалтын тоо хэмжээ', 'Орлуулсан борлуулалтын үнийн дүн',
    'Давуулсан борлуулалтын тоо хэмжээ', 'Давуулсан борлуулалтын үнийн дүн',
]


@app.route('/')
def index():
    # Эх файл солигдсон бол өмнөх шинжилгээ хуучирсан тул цэвэрлэнэ.
    if load_defaults():
        LAST_ANALYSIS.update(EMPTY_ANALYSIS)
    query = request.args.get('q', '')
    return render_template(
        'index.html',
        query=query,
        matches=search_substitutions(query),
        groups=STATE['groups'],
        products=STATE['substitutions'],
        logo_exists=logo_exists(),
        substitutions_loaded=bool(STATE['substitutions']),
        coverage=LAST_ANALYSIS['coverage'],
        matched=LAST_ANALYSIS['matched'],
        required_total=LAST_ANALYSIS['required_total'],
        lost_units_total=LAST_ANALYSIS['lost_units_total'],
        lost_revenue_total=LAST_ANALYSIS['lost_revenue_total'],
        outage_total=LAST_ANALYSIS['outage_total'],
        analysis=LAST_ANALYSIS['analysis'],
    )


@app.route('/analyze', methods=['POST'])
def analyze_route():
    # Боловсруулалт бүрийн өмнө эх файлуудын хамгийн сүүлийн хувилбарыг авна.
    load_defaults()

    def file_rows(field):
        uploaded = request.files.get(field)
        return rows_from_workbook(uploaded.read()) if uploaded else []

    sales_rows = file_rows('sales')
    previous_rows = file_rows('previous')
    warehouse_rows = file_rows('warehouse')
    outage_rows = file_rows('outage')

    outcome = analyze(sales_rows, previous_rows, warehouse_rows, outage_rows)
    results = outcome['analysis']
    LAST_ANALYSIS.update(outcome)
    if not results:
        flash('Тасалдлын тайлангаас боловсруулах мөр олдсонгүй.')
    return redirect(url_for('index'))


@app.route('/clear', methods=['POST'])
def clear_route():
    LAST_ANALYSIS.update(EMPTY_ANALYSIS)
    return redirect(url_for('index'))


@app.route('/download')
def download_route():
    results = LAST_ANALYSIS['analysis']
    if not results:
        flash('Эхлээд шинжилгээ хийнэ үү.')
        return redirect(url_for('index'))

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = 'Шинжилгээ'
    sheet.append(DOWNLOAD_HEADER)
    for row in results:
        sheet.append([
            row['code'],
            row['name'],
            row['lostUnits'],
            row['lostRevenue'],
            '; '.join(f"{item['code']} {item['name']}" for item in row['alternatives']),
            '; '.join(f"{item['code']}: {round(item['units'])} ш" for item in row['alternatives']),
            '; '.join(f"{item['code']}: {money(item['revenue'])}" for item in row['alternatives']),
            round(row['excess']),
            row['excessRevenue'],
        ])

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return send_file(
        buffer,
        as_attachment=True,
        download_name='SUBG_боловсруулалт.xlsx',
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )


load_defaults()

if __name__ == '__main__':
    # Байршуулсан орчинд 0.0.0.0:8000 (Dokploy-н Container Port) дээр сонсоно.
    app.run(
        host=os.environ.get('HOST', '0.0.0.0'),
        port=int(os.environ.get('PORT', '8000')),
        debug=os.environ.get('DEBUG', '').lower() in ('1', 'true', 'yes'),
    )
